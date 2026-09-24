"""pdf.tables — tables found by pdfplumber, as an XLSX workbook (one sheet per table) or CSVs."""

from __future__ import annotations

import csv
from pathlib import Path

import pdfplumber
from openpyxl import Workbook

from werkbank_engine.files import output_name
from werkbank_engine.jobs import Reporter, ToolContext, ToolError
from werkbank_engine.tools._pdf import parse_pages, pdf_inputs, zip_outputs

HEAVY = True
Table = list[list[str]]
NO_TABLES = "no tables found. Scanned pages need OCR first; tables without ruled lines are often missed."


def find_tables(report: Reporter, src: Path, name: str, spec: str) -> list[tuple[int, Table]]:
    try:
        pdf = pdfplumber.open(src)
    except Exception as exc:
        raise ToolError(
            f"{name} could not be read ({exc}). If it is protected, use Unlock PDF first."
        ) from None
    with pdf:
        pages = parse_pages(spec, len(pdf.pages))
        found = []
        for n, index in enumerate(pages):
            report(n / len(pages), f"Looking for tables on page {index + 1}")
            page = pdf.pages[index]
            for table in page.extract_tables():
                rows = [["" if cell is None else str(cell).strip() for cell in row] for row in table]
                if any(any(cell for cell in row) for row in rows):
                    found.append((index, rows))
            page.close()  # pdfplumber keeps parsed pages in memory otherwise
    return found


def write_xlsx(tables: list[tuple[int, Table]], dst: Path) -> None:
    book = Workbook()
    book.remove(book.active)
    counts: dict[int, int] = {}
    for index, rows in tables:
        counts[index] = counts.get(index, 0) + 1
        sheet = book.create_sheet(f"Page {index + 1} table {counts[index]}"[:31])
        for row in rows:
            sheet.append(row)
    book.save(dst)


def write_csvs(tables: list[tuple[int, Table]], folder: Path, stem: str, dst: Path) -> None:
    files = []
    counts: dict[int, int] = {}
    for index, rows in tables:
        counts[index] = counts.get(index, 0) + 1
        path = folder / f"{index}-{counts[index]}.csv"
        with path.open("w", encoding="utf-8-sig", newline="") as fh:  # BOM: Excel reads UTF-8 correctly
            csv.writer(fh).writerows(rows)
        files.append((path, f"{stem} p{index + 1} table {counts[index]}.csv"))
    zip_outputs(files, dst)


async def run(ctx: ToolContext) -> None:
    fmt, spec = ctx.params["format"], ctx.params["pages"]
    items = pdf_inputs(ctx)
    for i, name, src in items:
        ctx.progress(i / len(items), f"Reading {name}")
        tables = await ctx.in_thread(find_tables, src, name, spec)
        if not tables:
            ctx.note(f"{name}: {NO_TABLES}")
            continue
        if fmt == "xlsx":
            dst = ctx.work / f"{i}.xlsx"
            await ctx.in_thread(lambda report, t=tables, d=dst: write_xlsx(t, d))
            ctx.add_output(dst, output_name(name, "tables", ".xlsx"))
        else:
            folder = ctx.work / str(i)
            folder.mkdir()
            dst = ctx.work / f"{i}.zip"
            await ctx.in_thread(
                lambda report, t=tables, f=folder, d=dst, s=Path(name).stem: write_csvs(t, f, s, d)
            )
            ctx.add_output(dst, output_name(name, "tables", ".zip"))
        ctx.note(f"{name}: {len(tables)} table(s) found. Check them: table detection is not perfect.")
