"""The PDF tools, run through the API with the real libraries (pikepdf, PDFium, reportlab,
pdfplumber, img2pdf, pyHanko) and, for OCR, the real Tesseract."""

from __future__ import annotations

import datetime
import io
import json
import shutil
import zipfile
from pathlib import Path

import pikepdf
import pypdfium2 as pdfium
import pytest
from openpyxl import load_workbook
from PIL import Image, ImageDraw
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen.canvas import Canvas
from starlette.testclient import TestClient

from tests.conftest import run_job, upload
from werkbank_engine.config import Settings
from werkbank_engine.tools._pdf import PageSpecError, parse_pages
from werkbank_engine.tools.pdf_crop import unrotated_margins
from werkbank_engine.tools.pdf_layout import booklet_order

HAS_TESSERACT = shutil.which("tesseract") is not None


# --- fixtures -----------------------------------------------------------------------------------


def photo(width: int = 2600, height: int = 1800) -> Image.Image:
    """A noisy gradient: compresses like a photograph, not like flat colour."""
    image = Image.effect_noise((width, height), 60).convert("RGB")
    overlay = Image.linear_gradient("L").resize((width, height)).convert("RGB")
    return Image.blend(image, overlay, 0.5)


def make_sample(path: Path) -> Path:
    """Three A4 pages: text, a ruled table, a large photo; with metadata and an attachment."""
    canvas = Canvas(str(path), pagesize=A4)
    canvas.setTitle("Secret report")
    canvas.setAuthor("Jane Doe")
    canvas.setFont("Helvetica", 14)
    canvas.drawString(72, 760, "Page one: the quick brown fox. ID 8001015009087 belongs to Jane Doe.")
    canvas.showPage()
    canvas.setFont("Helvetica", 12)
    canvas.drawString(72, 760, "Page two has a table")
    rows = [["Name", "Mark"], ["Anna", "71"], ["Ben", "64"]]
    x0, y0, cw, rh = 72, 700, 150, 24
    for r, row in enumerate(rows):
        for c, cell in enumerate(row):
            canvas.drawString(x0 + c * cw + 6, y0 - r * rh - 16, cell)
    for r in range(len(rows) + 1):
        canvas.line(x0, y0 - r * rh, x0 + 2 * cw, y0 - r * rh)
    for c in range(3):
        canvas.line(x0 + c * cw, y0, x0 + c * cw, y0 - len(rows) * rh)
    canvas.showPage()
    buffer = io.BytesIO()
    photo().save(buffer, "JPEG", quality=95)
    from reportlab.lib.utils import ImageReader

    canvas.drawString(72, 800, "Page three has a photo")
    canvas.drawImage(ImageReader(io.BytesIO(buffer.getvalue())), 36, 200, width=520, height=360)
    canvas.showPage()
    canvas.save()
    with pikepdf.open(path, allow_overwriting_input=True) as pdf:
        pdf.attachments["notes.txt"] = pikepdf.AttachedFileSpec(pdf, b"attached notes")
        pdf.save(path)
    return path


@pytest.fixture
def sample(tmp_path: Path) -> Path:
    return make_sample(tmp_path / "report.pdf")


def text_of(path: Path) -> list[str]:
    doc = pdfium.PdfDocument(path)
    try:
        return [page.get_textpage().get_text_bounded() for page in doc]
    finally:
        doc.close()


def run(client: TestClient, tool: str, files: list[Path], params: dict | None = None) -> dict:
    job = run_job(client, tool, [{"fileId": upload(client, f)} for f in files], params)
    return job


def done(client: TestClient, tool: str, files: list[Path], params: dict | None = None) -> dict:
    job = run(client, tool, files, params)
    assert job["status"] == "done", job
    return job


def out(settings: Settings, name: str) -> Path:
    path = settings.outbox / name
    assert path.is_file(), sorted(p.name for p in settings.outbox.iterdir())
    return path


# --- page selections ----------------------------------------------------------------------------


def test_parse_pages() -> None:
    assert parse_pages("", 5) == [0, 1, 2, 3, 4]
    assert parse_pages("1-3, 5", 5) == [0, 1, 2, 4]
    assert parse_pages("4-2", 5) == [3, 2, 1]
    assert parse_pages("3-", 5) == [2, 3, 4]
    assert parse_pages("-2; 5", 5) == [0, 1, 4]
    for bad in ("0", "6", "a", "1-x", "--"):
        with pytest.raises(PageSpecError):
            parse_pages(bad, 5)
    with pytest.raises(PageSpecError):
        parse_pages("", 5, empty_means_all=False)


def test_booklet_order_and_crop_mapping() -> None:
    assert booklet_order(8) == [7, 0, 1, 6, 5, 2, 3, 4]
    assert booklet_order(3) == [None, 0, 1, 2]
    # Display margins (top, bottom, left, right) → unrotated (left, bottom, right, top).
    assert unrotated_margins(0, 1, 2, 3, 4) == (3, 2, 4, 1)
    assert unrotated_margins(90, 1, 2, 3, 4) == (1, 3, 2, 4)
    assert unrotated_margins(180, 1, 2, 3, 4) == (4, 1, 3, 2)
    assert unrotated_margins(270, 1, 2, 3, 4) == (2, 4, 1, 3)


# --- structure ----------------------------------------------------------------------------------


def test_merge_with_bookmarks(
    real_client: TestClient, settings: Settings, sample: Path, tmp_path: Path
) -> None:
    other = make_sample(tmp_path / "second.pdf")
    job = done(real_client, "pdf.merge", [sample, other])
    assert "6 pages" in job["notes"][0]
    with pikepdf.open(out(settings, "report (merged).pdf")) as pdf, pdf.open_outline() as outline:
        assert len(pdf.pages) == 6
        assert [item.title for item in outline.root] == ["report", "second"]


def test_split_modes(real_client: TestClient, settings: Settings, sample: Path) -> None:
    done(real_client, "pdf.split", [sample])
    with zipfile.ZipFile(out(settings, "report (split).zip")) as zf:
        assert sorted(zf.namelist()) == ["report p1.pdf", "report p2.pdf", "report p3.pdf"]
    done(real_client, "pdf.split", [sample], {"mode": "ranges", "ranges": "1-2, 3"})
    with zipfile.ZipFile(out(settings, "report (split) (2).zip")) as zf:
        assert sorted(zf.namelist()) == ["report p1-2.pdf", "report p3.pdf"]
    bad = run(real_client, "pdf.split", [sample], {"mode": "ranges", "ranges": "2-9"})
    assert bad["status"] == "failed" and "no page 9" in bad["error"]


def test_select_reorder_and_delete_pages(real_client: TestClient, settings: Settings, sample: Path) -> None:
    done(real_client, "pdf.pages", [sample], {"pages": "3, 1"})
    texts = text_of(out(settings, "report (pages).pdf"))
    assert len(texts) == 2 and "photo" in texts[0] and "Page one" in texts[1]
    done(real_client, "pdf.pages", [sample], {"action": "delete", "pages": "2"})
    assert len(text_of(out(settings, "report (pages) (2).pdf"))) == 2
    assert (
        "every page" in run(real_client, "pdf.pages", [sample], {"action": "delete", "pages": "1-3"})["error"]
    )
    assert "more than once" in run(real_client, "pdf.pages", [sample], {"pages": "1, 1"})["error"]


def test_rotate_selected_pages(real_client: TestClient, settings: Settings, sample: Path) -> None:
    done(real_client, "pdf.rotate", [sample], {"angle": "270", "pages": "2"})
    with pikepdf.open(out(settings, "report (rotated).pdf")) as pdf:
        assert [p.rotation for p in pdf.pages] == [0, 270, 0]


def test_layout_two_up_booklet_and_resize(real_client: TestClient, settings: Settings, sample: Path) -> None:
    done(real_client, "pdf.layout", [sample], {"layout": "2"})
    with pikepdf.open(out(settings, "report (2 per sheet).pdf")) as pdf:
        assert len(pdf.pages) == 2
        width, height = (float(v) for v in pdf.pages[0].mediabox[2:])
        assert width > height  # A4 landscape
    assert "Page one" in text_of(settings.outbox / "report (2 per sheet).pdf")[0]
    job = done(real_client, "pdf.layout", [sample], {"layout": "booklet"})
    assert "Print double-sided" in job["notes"][0]
    done(real_client, "pdf.layout", [sample], {"layout": "1", "paper": "letter"})
    with pikepdf.open(out(settings, "report (resized).pdf")) as pdf:
        assert [round(float(v)) for v in pdf.pages[0].mediabox[2:]] == [612, 792]


def test_crop_as_displayed(real_client: TestClient, settings: Settings, sample: Path) -> None:
    done(real_client, "pdf.crop", [sample], {"top": 20, "bottom": 0, "left": 0, "right": 0, "pages": "1"})
    with pikepdf.open(out(settings, "report (cropped).pdf")) as pdf:
        box = [float(v) for v in pdf.pages[0].cropbox]
        assert box[3] == pytest.approx(A4[1] - 20 * 72 / 25.4)
        assert pdf.pages[1].cropbox == pdf.pages[1].mediabox
    too_much = run(real_client, "pdf.crop", [sample], {"top": 300, "bottom": 0, "left": 0, "right": 0})
    assert "larger than page" in too_much["error"]


# --- security and clean-up ----------------------------------------------------------------------


def test_protect_with_aes256(real_client: TestClient, settings: Settings, sample: Path) -> None:
    job = done(real_client, "pdf.protect", [sample], {"password": "correct horse", "allowPrint": False})
    path = out(settings, "report (protected).pdf")
    with pytest.raises(pikepdf.PasswordError):
        pikepdf.open(path)
    with pikepdf.open(path, password="correct horse") as pdf:
        assert pdf.encryption.R == 6 and pdf.encryption.stream_method.name == "aesv3"
        assert not pdf.allow.print_highres and not pdf.allow.extract
    assert "correct horse" not in json.dumps([job, real_client.get("/api/jobs").json()])
    assert "at least 4" in run(real_client, "pdf.protect", [sample], {"password": "abc"})["error"]
    locked = run(real_client, "pdf.merge", [path, sample])
    assert locked["status"] == "failed" and "Unlock PDF" in locked["error"]


def test_remove_metadata(real_client: TestClient, settings: Settings, sample: Path) -> None:
    job = done(real_client, "pdf.clean-metadata", [sample])
    assert "Author" in job["notes"][0]
    with pikepdf.open(out(settings, "report (no metadata).pdf")) as pdf:
        assert "/Info" not in pdf.trailer or not pdf.trailer.Info.keys()
        assert "/Metadata" not in pdf.Root


def test_redaction_removes_the_text(real_client: TestClient, settings: Settings, sample: Path) -> None:
    job = done(real_client, "pdf.redact", [sample], {"terms": "8001015009087; jane doe"})
    assert "page(s) 1" in job["notes"][0]
    texts = text_of(out(settings, "report (redacted).pdf"))
    assert len(texts) == 3
    assert "8001015009087" not in texts[0] and "Jane" not in texts[0]
    assert "Page two has a table" in texts[1]  # pages without a match keep their text
    none = done(real_client, "pdf.redact", [sample], {"terms": "not in there"})
    assert "nothing was changed" in none["notes"][0] and none["outputs"] == []


def self_signed_p12(path: Path, password: str) -> Path:
    from cryptography import x509
    from cryptography.hazmat.primitives import hashes, serialization
    from cryptography.hazmat.primitives.asymmetric import rsa
    from cryptography.hazmat.primitives.serialization import pkcs12
    from cryptography.x509.oid import NameOID

    key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    name = x509.Name([x509.NameAttribute(NameOID.COMMON_NAME, "Test Signer")])
    now = datetime.datetime.now(datetime.UTC)
    cert = (
        x509.CertificateBuilder()
        .subject_name(name)
        .issuer_name(name)
        .public_key(key.public_key())
        .serial_number(x509.random_serial_number())
        .not_valid_before(now - datetime.timedelta(days=1))
        .not_valid_after(now + datetime.timedelta(days=30))
        .add_extension(
            x509.KeyUsage(True, True, False, False, False, False, False, False, False), critical=True
        )
        .sign(key, hashes.SHA256())
    )
    data = pkcs12.serialize_key_and_certificates(
        b"test", key, cert, None, serialization.BestAvailableEncryption(password.encode())
    )
    path.write_bytes(data)
    return path


def test_digital_signature(real_client: TestClient, settings: Settings, sample: Path, tmp_path: Path) -> None:
    from pyhanko.pdf_utils.reader import PdfFileReader
    from pyhanko.sign.validation import validate_pdf_signature
    from pyhanko_certvalidator import ValidationContext

    cert = self_signed_p12(tmp_path / "me.p12", "cert-pass")
    job = done(real_client, "pdf.sign", [sample, cert], {"password": "cert-pass", "reason": "Approved"})
    assert "Test Signer" in job["notes"][0]
    with out(settings, "report (signed).pdf").open("rb") as fh:
        signatures = PdfFileReader(fh).embedded_signatures
        assert len(signatures) == 1
        status = validate_pdf_signature(signatures[0], ValidationContext(trust_roots=[]))
        assert status.intact and status.valid  # untrusted (self-signed), but cryptographically sound
    wrong = run(real_client, "pdf.sign", [sample, cert], {"password": "nope"})
    assert wrong["status"] == "failed" and "certificate password" in wrong["error"]
    assert "cert-pass" not in json.dumps(real_client.get("/api/jobs").json())


def test_flatten_form(real_client: TestClient, settings: Settings, tmp_path: Path) -> None:
    path = tmp_path / "form.pdf"
    canvas = Canvas(str(path), pagesize=A4)
    canvas.acroForm.textfield(name="name", value="Filled In Value", x=72, y=700, width=300, height=24)
    canvas.showPage()
    canvas.save()
    job = done(real_client, "pdf.flatten", [path])
    assert "flattened 1" in job["notes"][0]
    result = out(settings, "form (flattened).pdf")
    with pikepdf.open(result) as pdf:
        assert "/AcroForm" not in pdf.Root and "/Annots" not in pdf.pages[0].obj
    assert "Filled In Value" in text_of(result)[0]


def test_repair_damaged_pdf(
    real_client: TestClient, settings: Settings, sample: Path, tmp_path: Path
) -> None:
    damaged = tmp_path / "damaged.pdf"
    data = sample.read_bytes()
    damaged.write_bytes(data[: data.rindex(b"xref")])  # cross-reference table and trailer gone
    job = done(real_client, "pdf.repair", [damaged])
    assert "repaired" in job["notes"][0]
    with pikepdf.open(out(settings, "damaged (repaired).pdf")) as pdf:
        assert len(pdf.pages) == 3


# --- content ------------------------------------------------------------------------------------


def test_compress_shrinks_photos_and_never_grows(
    real_client: TestClient, settings: Settings, sample: Path
) -> None:
    job = done(real_client, "pdf.compress", [sample], {"level": "strong"})
    assert "image(s) re-compressed" in job["notes"][0]
    result = out(settings, "report (compressed).pdf")
    assert result.stat().st_size < sample.stat().st_size * 0.7
    with pikepdf.open(result) as pdf:
        image = next(iter(pdf.pages[2].get_images().values()))
        assert max(int(image.Width), int(image.Height)) <= 1200
    again = done(real_client, "pdf.compress", [result], {"level": "strong"})
    assert out(settings, "report (compressed) (compressed).pdf").stat().st_size <= result.stat().st_size
    assert again["notes"]


def test_page_numbers_on_a_rotated_page(real_client: TestClient, settings: Settings, sample: Path) -> None:
    with pikepdf.open(sample, allow_overwriting_input=True) as pdf:
        pdf.pages[1].rotate(90, relative=True)
        pdf.save(sample)
    done(real_client, "pdf.page-numbers", [sample], {"style": "page-n-of-total"})
    result = out(settings, "report (numbered).pdf")
    doc = pdfium.PdfDocument(result)
    for n, page in enumerate(doc):
        textpage = page.get_textpage()
        searcher = textpage.search(f"Page {n + 1} of 3")
        found = searcher.get_next()
        assert found is not None, n
        _, bottom, _, _ = textpage.get_charbox(found[0])
        _, height = page.get_size()
        assert bottom < height * 0.1, (n, bottom, height)  # at the bottom as displayed
    doc.close()


def test_watermark(real_client: TestClient, settings: Settings, sample: Path) -> None:
    done(real_client, "pdf.watermark", [sample], {"text": "DRAFT", "pages": "1-2"})
    texts = text_of(out(settings, "report (watermarked).pdf"))
    assert ["DRAFT" in t for t in texts] == [True, True, False]
    bad = run(real_client, "pdf.watermark", [sample], {"text": "漢字"})
    assert bad["status"] == "failed" and "Latin letters" in bad["error"]


def test_to_images_and_back(
    real_client: TestClient, settings: Settings, sample: Path, tmp_path: Path
) -> None:
    done(real_client, "pdf.to-images", [sample], {"dpi": 72})
    with zipfile.ZipFile(out(settings, "report (PNG images).zip")) as zf:
        assert sorted(zf.namelist()) == ["report p1.png", "report p2.png", "report p3.png"]
    done(real_client, "pdf.to-images", [sample], {"format": "jpg", "pages": "2"})
    with Image.open(out(settings, "report p2.jpg")) as image:
        assert image.size[0] == pytest.approx(A4[0] * 150 / 72, abs=1)
        assert image.size[1] == pytest.approx(A4[1] * 150 / 72, abs=1)

    jpeg = tmp_path / "photo.jpg"
    photo(800, 600).save(jpeg, "JPEG", quality=85)
    png = tmp_path / "logo.png"
    logo = Image.new("RGBA", (300, 300), (0, 0, 0, 0))
    ImageDraw.Draw(logo).ellipse((20, 20, 280, 280), fill=(200, 0, 0, 255))
    logo.save(png)
    job = done(real_client, "pdf.from-images", [jpeg, png])
    assert "2 embedded without re-compression" in job["notes"][0]  # img2pdf keeps PNG transparency
    with pikepdf.open(out(settings, "photo (images).pdf")) as pdf:
        assert len(pdf.pages) == 2
        assert [round(float(v)) for v in pdf.pages[0].mediabox[2:]] == [842, 595]  # landscape photo, A4
        embedded = next(iter(pdf.pages[0].get_images().values()))
        assert embedded.read_raw_bytes() == jpeg.read_bytes()  # the JPEG was not re-compressed


def test_text_tables_info_extract_compare(
    real_client: TestClient, settings: Settings, sample: Path, tmp_path: Path
) -> None:
    done(real_client, "pdf.to-text", [sample])
    text = out(settings, "report (text).txt").read_text(encoding="utf-8")
    assert "--- Page 1 ---" in text and "quick brown fox" in text

    done(real_client, "pdf.tables", [sample])
    book = load_workbook(out(settings, "report (tables).xlsx"))
    rows = [[c.value for c in row] for row in book.worksheets[0].iter_rows()]
    assert ["Anna", "71"] in rows and book.sheetnames == ["Page 2 table 1"]

    done(real_client, "pdf.info", [sample])
    info = out(settings, "report (info).txt").read_text(encoding="utf-8")
    for expected in (
        "Pages: 3",
        "A4 portrait (3)",
        "Author: Jane Doe",
        "Attachments: 1 (notes.txt)",
        "Images: 1",
    ):
        assert expected in info, info

    done(real_client, "pdf.extract", [sample])
    with zipfile.ZipFile(out(settings, "report (images).zip")) as zf:
        assert len(zf.namelist()) == 1 and zf.namelist()[0].endswith(".jpg")
    done(real_client, "pdf.extract", [sample], {"what": "attachments"})
    with zipfile.ZipFile(out(settings, "report (attachments).zip")) as zf:
        assert zf.read("notes.txt") == b"attached notes"

    changed = tmp_path / "changed.pdf"
    canvas = Canvas(str(changed), pagesize=A4)
    canvas.setFont("Helvetica", 14)
    canvas.drawString(72, 760, "Page one: the quick red fox. ID 8001015009087 belongs to Jane Doe.")
    canvas.save()
    job = done(real_client, "pdf.compare", [sample, changed])
    assert "changed line(s)" in job["notes"][0]
    html = out(settings, "report (comparison).html").read_text(encoding="utf-8")
    assert "diff_chg" in html and "quick" in html  # the changed line is highlighted


@pytest.mark.skipif(not HAS_TESSERACT, reason="tesseract not installed")
def test_ocr_makes_scans_searchable(
    real_client: TestClient, settings: Settings, tmp_path: Path, sample: Path
) -> None:
    # A "scan": page 1 of the sample rendered to an image, plus a page that already has text.
    doc = pdfium.PdfDocument(sample)
    image = doc[0].render(scale=300 / 72).to_pil().convert("L")
    doc.close()
    scan_png = tmp_path / "scan.png"
    image.save(scan_png, dpi=(300, 300))
    import img2pdf

    scan = tmp_path / "scan.pdf"
    scan.write_bytes(img2pdf.convert(str(scan_png)))
    with pikepdf.open(scan, allow_overwriting_input=True) as pdf, pikepdf.open(sample) as text_pdf:
        pdf.pages.append(text_pdf.pages[1])
        pdf.save(scan)
    assert "quick" not in text_of(scan)[0]

    job = done(real_client, "pdf.ocr", [scan])
    assert "1 page(s); 1 page(s) already had text" in job["notes"][0]
    result = out(settings, "scan (searchable).pdf")
    texts = text_of(result)
    assert "quick brown fox" in texts[0]
    assert "Page two has a table" in texts[1]
    with pikepdf.open(result) as pdf:  # the scan itself is untouched: same image, same size
        assert next(iter(pdf.pages[0].get_images().values())).Width == image.width
